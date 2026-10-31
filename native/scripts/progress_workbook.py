#!/usr/bin/env python3
"""Build the decompilation status workbook (notes/SC4_decompilation_status.xlsx).

Reads the same sources as scripts/progress.py (notes/symbols-functions-demangled.txt +
notes/ledger.csv) and writes one sheet per view: overview, binary structure, every game class,
and every function that is still missing. --ghidra-total overrides the Ghidra function count.

usage: scripts/progress_workbook.py [--ghidra-total 46143] [--out PATH]
"""
import argparse, collections, csv, datetime, pathlib, re, sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = pathlib.Path(__file__).resolve().parent.parent

# Reuse the classification and owner parsing of scripts/progress.py without running its main().
_src = (ROOT / "scripts/progress.py").read_text().split("def main()")[0]
_ns = {"__file__": str(ROOT / "scripts/progress.py")}
exec(_src, _ns)
classify, owner_of = _ns["classify"], _ns["owner_of"]

LAYER_NOTE = {
    "cSC4 game logic (cSC4*, SC4*, cSTE*, cSL*)": "City simulation, tools, views, window clients, terrain, properties, Lua game side",
    "cGZ framework": "EA's GZCOM framework: COM, windows, draw context, DBPF, resources, fonts, audio director",
    "other game classes": "Game classes outside the cSC4/cGZ prefixes (occupants, S3D, HTML, misc toolkits)",
    "cRZ framework": "Maxis' portable layer under GZCOM: strings, files, streams, variants, dates, random",
    "free C functions (unclassified)": "Unqualified C functions: helpers whose owning class is not in the symbol",
    "Audio (Hit*/Fresh*)": "HIT/Fresh audio middleware (sample playback, track data blocks)",
    "Graphics driver layer (DX7 shim, DrawContext, texels)": "Direct3D 7 shim, draw context, texel operations, SGL rasteriser",
    "stream operators etc. (operator>>/<<)": "Global stream operators",
    "Aspyr / Mac glue": "Aspyr's Mac port glue (MacTools, NCT_*)",
    "libc++ (std::__1)": "Linked third-party: libc++ — not rewritten, only stubbed where the game touches it",
    "C libs (FreeType, Lua, zlib/png, SDL/SND)": "Linked third-party: FreeType, Lua 5.0, zlib/libpng/libjpeg, SDL, SND/AEMS",
    "compiler thunks / static-init glue": "Compiler-generated: thunks, guards, static initialisers, operators new/delete",
}

HEAD = Font(bold=True, color="FFFFFF", size=11)
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
TITLE = Font(bold=True, size=16, color="1F3864")
SUB = Font(italic=True, size=10, color="595959")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def width(sheet, widths):
    for i, w in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = w


def header_row(sheet, row, values, wrap=True):
    for col, v in enumerate(values, start=1):
        c = sheet.cell(row=row, column=col, value=v)
        c.font = HEAD
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
        c.border = BOX


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ghidra-total", type=int, default=46143)
    ap.add_argument("--out", default=str(ROOT / "notes/SC4_decompilation_status.xlsx"))
    a = ap.parse_args()

    funcs = {}
    for line in (ROOT / "notes/symbols-functions-demangled.txt").read_text(errors="replace").splitlines():
        addr, _, name = line.partition(" ")
        funcs[int(addr, 16)] = name
    done = {}
    for r in csv.DictReader((ROOT / "notes/ledger.csv").open()):
        done[int(r["address"], 16)] = r["status"]

    totals = collections.Counter()
    rec = collections.Counter()
    classes = collections.Counter()
    class_layer = {}
    class_rec = collections.Counter()
    remaining = []
    for addr, name in funcs.items():
        kind, cat = classify(name, addr)
        totals[(kind, cat)] += 1
        recovered = done.get(addr, "").startswith("recovered")
        if recovered:
            rec[(kind, cat)] += 1
        if kind != "game":
            continue
        owner, _ = owner_of(name)
        key = owner or "(free functions)"
        classes[key] += 1
        class_layer.setdefault(key, cat)
        if key == "(free functions)":
            class_layer[key] = "free C functions (unclassified)"
        if recovered:
            class_rec[key] += 1
        else:
            remaining.append((cat, key, addr, done.get(addr, "not started"), name))

    nlib = sum(v for (k, _), v in totals.items() if k == "lib")
    ncomp = sum(v for (k, _), v in totals.items() if k == "compiler")
    ngame = sum(v for (k, _), v in totals.items() if k == "game")
    nrec = sum(v for (k, _), v in rec.items() if k == "game")
    pct = nrec / ngame

    wb = Workbook()

    # ---------------------------------------------------------------- Overview
    ov = wb.active
    ov.title = "Overview"
    ov["A1"] = "SimCity 4 (Aspyr macOS x86_64) — decompilation status"
    ov["A1"].font = TITLE
    ov["A2"] = ("Binary: Sim City 4 Deluxe Editionsub, v1.2.1 build 185949 · "
                f"generated {datetime.date.today().isoformat()} · source: notes/ledger.csv + Ghidra MCP")
    ov["A2"].font = SUB
    rows = [
        ("Functions in Ghidra", a.ghidra_total, "everything Ghidra has as a function, named or not"),
        ("With a demangled symbol", len(funcs), "named by the symbol table"),
        ("Unnamed / stub", a.ghidra_total - len(funcs), "thunks or functions without a symbol"),
        ("Third-party libraries", nlib, "linked, not rewritten"),
        ("Compiler thunks / static init", ncomp, "regenerated by the compiler"),
        ("Game code to recover", ngame, "the actual target"),
        ("Recovered", nrec, "hand-written and ledgered"),
        ("Remaining", ngame - nrec, "still to do"),
        ("Completion", pct, "recovered / game code"),
    ]
    header_row(ov, 4, ["Metric", "Value", "Note"])
    r = 5
    for label, value, note in rows:
        ov.cell(row=r, column=1, value=label).border = BOX
        c = ov.cell(row=r, column=2, value=value)
        c.border = BOX
        if label == "Completion":
            c.number_format = "0.0%"
            c.font = Font(bold=True)
        else:
            c.number_format = "#,##0"
        ov.cell(row=r, column=3, value=note).border = BOX
        r += 1
    width(ov, [34, 34, 46])

    # layer table
    start = 16
    ov.cell(row=start - 1, column=1, value="By layer").font = Font(bold=True, size=12)
    header_row(ov, start, ["Layer", "Functions", "Recovered", "Remaining", "% done"])
    r = start + 1
    game_layers = []
    for (kind, cat), v in sorted(totals.items(), key=lambda x: -x[1]):
        if kind != "game":
            continue
        game_layers.append((cat, v, rec[(kind, cat)]))
        ov.cell(row=r, column=1, value=cat).border = BOX
        ov.cell(row=r, column=2, value=v).border = BOX
        ov.cell(row=r, column=3, value=rec[(kind, cat)]).border = BOX
        ov.cell(row=r, column=4, value=v - rec[(kind, cat)]).border = BOX
        pc = ov.cell(row=r, column=5, value=rec[(kind, cat)] / v)
        pc.number_format = "0.0%"
        pc.border = BOX
        for col in (2, 3, 4):
            ov.cell(row=r, column=col).number_format = "#,##0"
        r += 1
    for (kind, cat), v in sorted(totals.items(), key=lambda x: -x[1]):
        if kind == "game":
            continue
        ov.cell(row=r, column=1, value=f"{cat} ({'library' if kind == 'lib' else 'compiler'})").border = BOX
        ov.cell(row=r, column=2, value=v).border = BOX
        ov.cell(row=r, column=3, value="—").border = BOX
        ov.cell(row=r, column=4, value="—").border = BOX
        ov.cell(row=r, column=5, value="—").border = BOX
        ov.cell(row=r, column=2).number_format = "#,##0"
        r += 1

    chart = BarChart()
    chart.type = "bar"
    chart.title = "Recovered vs remaining, by layer"
    chart.height, chart.width = 9, 18
    data = Reference(ov, min_col=3, max_col=4, min_row=start, max_row=start + len(game_layers))
    cats = Reference(ov, min_col=1, min_row=start + 1, max_row=start + len(game_layers))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.y_axis.title = "functions"
    ov.add_chart(chart, "G4")

    # ---------------------------------------------------------------- Structure
    st = wb.create_sheet("Structure")
    st["A1"] = "What the binary is made of"
    st["A1"].font = TITLE
    st["A2"] = "Layers, size and what each one is. Percentages are of the 43,280 symbolised functions."
    st["A2"].font = SUB
    header_row(st, 4, ["Layer", "Functions", "Share", "Recovered", "Remaining", "% done", "What it is"])
    r = 5
    for (kind, cat), v in sorted(totals.items(), key=lambda x: -x[1]):
        st.cell(row=r, column=1, value=cat).border = BOX
        st.cell(row=r, column=2, value=v).border = BOX
        st.cell(row=r, column=2).number_format = "#,##0"
        s = st.cell(row=r, column=3, value=v / len(funcs))
        s.number_format = "0.0%"
        s.border = BOX
        if kind == "game":
            st.cell(row=r, column=4, value=rec[(kind, cat)]).border = BOX
            st.cell(row=r, column=4).number_format = "#,##0"
            st.cell(row=r, column=5, value=v - rec[(kind, cat)]).border = BOX
            st.cell(row=r, column=5).number_format = "#,##0"
            p = st.cell(row=r, column=6, value=rec[(kind, cat)] / v)
        else:
            st.cell(row=r, column=4, value="—").border = BOX
            st.cell(row=r, column=5, value="—").border = BOX
            p = st.cell(row=r, column=6, value="—")
        if isinstance(p.value, float):
            p.number_format = "0.0%"
        p.border = BOX
        st.cell(row=r, column=7, value=LAYER_NOTE.get(cat, "")).border = BOX
        st.cell(row=r, column=7).alignment = Alignment(wrap_text=True, vertical="top")
        r += 1
    st.cell(row=r, column=1, value="TOTAL").font = Font(bold=True)
    st.cell(row=r, column=2, value=len(funcs)).font = Font(bold=True)
    st.cell(row=r, column=3, value=1.0).number_format = "0.0%"
    st.cell(row=r, column=4, value=nrec).number_format = "#,##0"
    st.cell(row=r, column=5, value=ngame - nrec).number_format = "#,##0"
    st.cell(row=r, column=6, value=pct).number_format = "0.0%"
    width(st, [46, 11, 8, 11, 11, 8, 70])

    pie = PieChart()
    pie.title = "Share of the 43,280 symbolised functions"
    labels = Reference(st, min_col=1, min_row=5, max_row=r - 1)
    data = Reference(st, min_col=2, min_row=4, max_row=r - 1)
    pie.add_data(data, titles_from_data=True)
    pie.set_categories(labels)
    pie.height, pie.width = 10, 14
    st.add_chart(pie, "I4")

    # ---------------------------------------------------------------- Classes
    cl = wb.create_sheet("Classes")
    cl["A1"] = "Every game class, biggest gaps first"
    cl["A1"].font = TITLE
    cl["A2"] = (f"{len(classes):,} classes / {ngame:,} game functions. "
                "Join the class name to src/<group>/<name>.{h,cpp}; the ledger names the recovered ones.")
    cl["A2"].font = SUB
    header_row(cl, 4, ["#", "Class", "Layer", "Functions", "Recovered", "Remaining", "% done"])
    order = sorted(classes.items(), key=lambda kv: (-(kv[1] - class_rec[kv[0]]), -kv[1]))
    r = 5
    for i, (name, v) in enumerate(order, start=1):
        cl.cell(row=r, column=1, value=i).border = BOX
        cl.cell(row=r, column=2, value=name).border = BOX
        cl.cell(row=r, column=3, value=class_layer[name]).border = BOX
        cl.cell(row=r, column=4, value=v).border = BOX
        cl.cell(row=r, column=4).number_format = "#,##0"
        cl.cell(row=r, column=5, value=class_rec[name]).border = BOX
        cl.cell(row=r, column=5).number_format = "#,##0"
        cl.cell(row=r, column=6, value=v - class_rec[name]).border = BOX
        cl.cell(row=r, column=6).number_format = "#,##0"
        p = cl.cell(row=r, column=7, value=class_rec[name] / v)
        p.number_format = "0.0%"
        p.border = BOX
        r += 1
    cl.freeze_panes = "C5"
    cl.auto_filter.ref = f"A4:G{r - 1}"
    cl.conditional_formatting.add(f"G5:G{r - 1}",
                                  ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                                                 mid_type="num", mid_value=0.5, mid_color="FFEB84",
                                                 end_type="num", end_value=1, end_color="63BE7B"))
    width(cl, [6, 46, 44, 11, 11, 11, 9])

    # ---------------------------------------------------------------- Remaining
    rm = wb.create_sheet("Remaining functions")
    rm["A1"] = f"Functions still to recover ({len(remaining):,})"
    rm["A1"].font = TITLE
    rm["A2"] = "Status comes from notes/ledger.csv: 'not started' or 'pending' (blocked on a dependency)."
    rm["A2"].font = SUB
    header_row(rm, 4, ["Demangled symbol", "Class", "Layer", "Status", "Address"])
    r = 5
    for cat, cls, addr, status, name in sorted(remaining, key=lambda x: (x[0], x[1], x[2])):
        rm.cell(row=r, column=1, value=name).border = BOX
        rm.cell(row=r, column=2, value=cls).border = BOX
        rm.cell(row=r, column=3, value=cat).border = BOX
        rm.cell(row=r, column=4, value=status).border = BOX
        c = rm.cell(row=r, column=5, value=f"0x{addr:x}")
        c.border = BOX
        r += 1
    rm.freeze_panes = "A5"
    rm.auto_filter.ref = f"A4:E{r - 1}"
    width(rm, [96, 42, 42, 14, 14])

    out = pathlib.Path(a.out)
    wb.save(out)
    print(f"wrote {out}")
    print(f"game functions {ngame:,}, recovered {nrec:,} ({pct:.2%}), remaining {ngame - nrec:,}; classes {len(classes):,}")


main()
